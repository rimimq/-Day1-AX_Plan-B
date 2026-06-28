#!/usr/bin/env python3
"""Build a Google Sheets-ready workbook for the Plan-B week 1 flow report."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "week1-google-sheets"
OUTPUT = ROOT / "week1-google-sheets" / "Plan-B_1주차_전체플로우.xlsx"

SHEETS = [
    ("01_요약", SOURCE_DIR / "01_요약.csv"),
    ("02_전체플로우", SOURCE_DIR / "02_전체_플로우.csv"),
    ("03_시트구조", SOURCE_DIR / "03_시트_구조.csv"),
    ("04_단계별상세", SOURCE_DIR / "04_단계별_상세.csv"),
    ("05_확인질문", SOURCE_DIR / "05_확인_질문.csv"),
]


def read_csv(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f))


def set_column_widths(ws) -> None:
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in column_cells:
            text = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(text), 48))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(max_len + 3, 48))


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = body_alignment

    for row_idx in range(1, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 24 if row_idx == 1 else 42

    set_column_widths(ws)


def main() -> None:
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_name, csv_path in SHEETS:
        ws = wb.create_sheet(sheet_name)
        for row in read_csv(csv_path):
            ws.append(row)
        style_sheet(ws)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
